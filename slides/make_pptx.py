"""Dựng bản PowerPoint của bộ slide: mỗi trang main.pdf thành một ảnh tràn khung, kèm ghi chú
thuyết trình lấy từ SPEAKER_NOTES.md (mục "### Slide N — ...") để xem ở Presenter View.

Chạy từ thư mục slides/, sau khi build main.pdf:  python make_pptx.py
Cần python-pptx và PyMuPDF (có ở python base của anaconda).
"""
import io
import re
import sys
from pathlib import Path

import fitz
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
PDF = HERE / 'main.pdf'
NOTES = HERE / 'SPEAKER_NOTES.md'
OUT = HERE / 'Vietnamese-DBpedia-slides.pptx'
DPI = 300


def speaker_notes():
    """{số slide: văn bản ghi chú} từ các mục '### Slide N — tiêu đề'."""
    text = NOTES.read_text(encoding='utf-8')
    notes = {}
    for m in re.finditer(r'^### Slide (\d+)[^\n]*\n(.*?)(?=^### |^## |^---|\Z)', text, re.M | re.S):
        # Mỗi nhãn (**Nói:**, **Số phải trúng:**, **Nếu bị hỏi** ...) mở một đoạn mới;
        # các dòng còn lại nối vào đoạn đang viết.
        paras = []
        for line in m.group(2).splitlines():
            s = line.strip()
            if not s:
                continue
            if not paras or s.startswith('**') or s.startswith('*('):
                paras.append(s)
            else:
                paras[-1] += ' ' + s
        body = '\n'.join(paras)
        body = re.sub(r'\*\*(.*?)\*\*', r'\1', body)          # **đậm**
        body = re.sub(r'(?<!\w)\*(.*?)\*(?!\w)', r'\1', body)  # *nghiêng*
        notes[int(m.group(1))] = body.replace('`', '')
    return notes


def main():
    pdf = fitz.open(PDF)
    notes = speaker_notes()
    prs = Presentation()
    w_pt, h_pt = pdf[0].rect.width, pdf[0].rect.height
    prs.slide_width = Emu(9144000)                         # 10 in, 4:3 như mẫu HUST
    prs.slide_height = Emu(round(9144000 * h_pt / w_pt))
    blank = prs.slide_layouts[6]
    for i, page in enumerate(pdf, start=1):
        png = page.get_pixmap(dpi=DPI).tobytes('png')
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(io.BytesIO(png), 0, 0, prs.slide_width, prs.slide_height)
        if i in notes:
            slide.notes_slide.notes_text_frame.text = notes[i]
    prs.save(OUT)
    missing = [i for i in range(1, len(pdf) + 1) if i not in notes]
    print(f'wrote {OUT.name}: {len(pdf)} slides, notes on {len(pdf) - len(missing)}'
          + (f' (no notes: {missing})' if missing else ''))
    extra = sorted(set(notes) - set(range(1, len(pdf) + 1)))
    if extra:
        print(f'WARNING: SPEAKER_NOTES.md has notes for slides that do not exist: {extra}')
        sys.exit(1)


if __name__ == '__main__':
    main()
